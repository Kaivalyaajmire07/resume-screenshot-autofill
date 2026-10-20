"""Export an extraction result as a formatted Word resume (python-docx, fully local)."""
from __future__ import annotations

from io import BytesIO


def _fmt_exp(e) -> tuple[str, str, str]:
    role = e.title
    if role and not role.lower().startswith(("worked", "working")):
        role = f"Worked as a {role}"
    end = "Present" if e.current else e.end_date
    dates = " – ".join(x for x in (e.start_date, end) if x)
    return e.company, role, "  |  ".join(x for x in (dates, e.duration) if x)


def build_docx(values: dict, education: list, experience: list) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    def para(text, bold=False, size=None, align=None, after=2, color=None):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor(*color)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        return p

    def head(text):
        p = para(text.upper(), True, 12, after=4, color=(0x1F, 0x4E, 0x79))
        p.paragraph_format.space_before = Pt(10)
        border, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F4E79")):
            bottom.set(qn("w:" + k), v)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)

    centre = WD_ALIGN_PARAGRAPH.CENTER
    name = " ".join(x for x in (values.get("first_name"), values.get("middle_name"), values.get("last_name")) if x)
    para(name or "Candidate", True, 22, centre)
    if values.get("address"):
        para(values["address"], size=10, align=centre, after=1)
    para("  |  ".join(x for x in (values.get("email"), values.get("mobile")) if x), size=10, align=centre, after=4)

    for title, key in (("Objective", "objectives"), ("Summary", "summary")):
        if values.get(key):
            head(title)
            para(values[key])

    head("Professional Experience")
    if experience:
        for e in experience:
            company, role, meta = _fmt_exp(e)
            para(company or role, True, after=0)
            if company and role:
                para(role, after=0)
            para(meta, after=6)
    else:
        para("No professional experience found.")

    head("Education")
    if education:
        for r in education:
            para(r.institution or (r.degree or r.level.replace("_", " ").title()), True, after=0)
            bits = [r.degree if r.institution else "", r.specialization, r.board_or_body, r.status if r.status != "Not Available" else "",
                    r.passing_year or (f"Expected {r.expected_year}" if r.expected_year else ""),
                    f"{r.percentage}" if r.percentage else (f"CGPA {r.cgpa}" if r.cgpa else "")]
            line = "  |  ".join(b for b in bits if b)
            para((r.level.replace("_", " ").title() + (": " + line if line else "")), after=6)
    else:
        para("No education found.")

    skills = [x.lstrip("•-* ").strip() for x in (values.get("key_skills") or "").splitlines() if x.strip()]
    if skills:
        head("Key Skills")
        for k in skills:
            d.add_paragraph(k, style="List Bullet").paragraph_format.space_after = Pt(0)

    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()
